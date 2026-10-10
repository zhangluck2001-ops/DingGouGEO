# ========= Copyright 2025-2026 @ Eigent.ai All Rights Reserved. =========
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.
# ========= Copyright 2025-2026 @ Eigent.ai All Rights Reserved. =========
"""Turn uploaded brand materials into text the cleaning agent can read.

Parsing is done once, by code, when a file is uploaded (or before a cleaning
task if it was dropped into the folder by hand) — not by the agent inside its
loop, where it improvised OCR programs, took minutes and left build caches in
the knowledge base. For each file in ``原始资料/`` this writes
``原始资料解析/<文件名>/``:

- ``内容.md`` — the text, with every part labelled by where it is in the
  original (page, sheet!cell, embedded attachment), so facts can cite it;
- ``图片/`` — images worth looking at (embedded pictures, scanned pages),
  downscaled; logos and icons are skipped;
- ``清单.json`` — counts, which images need looking at, attachments, errors.

Handles PDF (text, tables, scanned pages), Excel (sheets, pictures placed
over or in cells — including WPS ``DISPIMG`` cell images — and files embedded
as OLE packages, parsed in turn), Word/PowerPoint (text and pictures),
images, and plain text. Anything else goes through MarkItDown.
"""

import hashlib
import io
import json
import logging
import posixpath
import re
import shutil
import struct
import unicodedata
import zipfile
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from defusedxml import ElementTree

logger = logging.getLogger("geo.source_parse")

PARSE_FOLDER = "原始资料解析"
CONTENT_FILE = "内容.md"
MANIFEST_FILE = "清单.json"
IMAGE_FOLDER = "图片"
ATTACHMENT_FOLDER = "附件"
# Bump when the output changes so earlier results are parsed again.
PARSER_VERSION = 2

IMAGE_SUFFIXES = {
    ".png",
    ".jpg",
    ".jpeg",
    ".webp",
    ".gif",
    ".bmp",
    ".tif",
    ".tiff",
}
TEXT_SUFFIXES = {".md", ".txt", ".csv", ".json"}
# A PDF page with less text than this is treated as scanned and rendered.
SCANNED_PAGE_CHARS = 40
# Pictures smaller than this (bytes or pixels) are logos and icons.
MIN_IMAGE_BYTES = 8 * 1024
MIN_IMAGE_SIDE = 120
MAX_IMAGE_SIDE = 2000
RENDER_SCALE = 2.0  # 144 dpi
MAX_ATTACHMENT_DEPTH = 1

_NS = {
    "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
}
_DISPIMG = re.compile(r'DISPIMG\(\s*"([^"]+)"')
# Kangxi and CJK radical code points PDFs often emit instead of the
# characters they look like (「⽜」 for 「牛」), which breaks keyword matching.
_RADICALS = re.compile("[\u2e80-\u2eff\u2f00-\u2fdf]")


def _fix_radicals(text: str) -> str:
    return _RADICALS.sub(
        lambda m: unicodedata.normalize("NFKC", m.group()), text
    )


@dataclass
class _Result:
    """What one file (and its attachments) turned into."""

    out: Path
    lines: list[str] = field(default_factory=list)
    images: list[dict] = field(default_factory=list)
    attachments: list[dict] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    pages: int = 0
    sheets: int = 0
    tables: int = 0
    _seen: set[str] = field(default_factory=set)

    def add_image(self, data: bytes, name: str, where: str) -> None:
        """Keep an image worth looking at; logos, icons and repeats are skipped."""
        digest = hashlib.sha1(data, usedforsecurity=False).hexdigest()[:12]
        entry: dict = {"where": where, "bytes": len(data)}
        if digest in self._seen:
            return
        self._seen.add(digest)
        suffix = Path(name).suffix.lower()
        if suffix in {".emf", ".wmf"}:
            entry.update(look=False, reason="图标（emf/wmf）")
            self.images.append(entry)
            return
        try:
            from PIL import Image

            with Image.open(io.BytesIO(data)) as img:
                img.load()
                width, height = img.size
                entry["size"] = [width, height]
                if (
                    len(data) < MIN_IMAGE_BYTES
                    or min(width, height) < MIN_IMAGE_SIDE
                ):
                    entry.update(look=False, reason="太小，按图标或装饰处理")
                    self.images.append(entry)
                    return
                if max(width, height) > MAX_IMAGE_SIDE:
                    img.thumbnail((MAX_IMAGE_SIDE, MAX_IMAGE_SIDE))
                rgb = (
                    img.convert("RGB") if img.mode not in ("RGB", "L") else img
                )
                target = (
                    self.out
                    / IMAGE_FOLDER
                    / f"{_safe(Path(name).stem)}-{digest}.jpg"
                )
                target.parent.mkdir(parents=True, exist_ok=True)
                rgb.save(target, "JPEG", quality=85)
        except Exception as exc:  # noqa: BLE001 — keep the rest of the file
            entry.update(look=False, reason=f"读不了这张图：{exc}")
            self.images.append(entry)
            return
        entry.update(file=str(target.relative_to(self.out)), look=True)
        self.images.append(entry)


def _safe(name: str) -> str:
    return re.sub(r"[^\w一-鿿.-]+", "_", name).strip("_")[:60] or "file"


def _md_table(rows: list[list]) -> str:
    rows = [
        [
            "" if c is None else str(c).replace("\n", " ").replace("|", "\\|")
            for c in r
        ]
        for r in rows
    ]
    rows = [r for r in rows if any(cell.strip() for cell in r)]
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    head, *body = rows
    out = ["| " + " | ".join(head) + " |", "|" + "---|" * width]
    out += ["| " + " | ".join(r) + " |" for r in body]
    return "\n".join(out)


# ----------------------------------------------------------------- PDF


def _parse_pdf(data: bytes, label: str, res: _Result) -> None:
    import pdfplumber

    with pdfplumber.open(io.BytesIO(data)) as pdf:
        res.pages += len(pdf.pages)
        scanned: list[int] = []
        for number, page in enumerate(pdf.pages, start=1):
            res.lines.append(f"\n### {label} 第 {number} 页\n")
            tables = page.find_tables()
            rest = page
            for table in tables:
                try:
                    rest = rest.outside_bbox(table.bbox)
                except ValueError:
                    pass
            text = (rest.extract_text() or "").strip()
            if text:
                res.lines.append(text)
            for table in tables:
                md = _md_table(table.extract())
                if md:
                    res.tables += 1
                    res.lines.append(f"\n{md}\n")
            all_text = (page.extract_text() or "").strip()
            if len(all_text) < SCANNED_PAGE_CHARS:
                scanned.append(number)
                res.lines.append(
                    "（这一页基本没有文字层，可能是扫描件，见「需要看图」）"
                )
    if scanned:
        _render_pages(data, label, scanned, res)


def _render_pages(
    data: bytes, label: str, numbers: list[int], res: _Result
) -> None:
    try:
        import pypdfium2 as pdfium
    except ImportError:
        res.errors.append(f"{label}：缺少 pypdfium2，扫描页无法转成图片")
        return
    doc = pdfium.PdfDocument(data)
    try:
        for number in numbers:
            image = doc[number - 1].render(scale=RENDER_SCALE).to_pil()
            buffer = io.BytesIO()
            image.convert("RGB").save(buffer, "JPEG", quality=85)
            res.add_image(
                buffer.getvalue(),
                f"{label}-p{number}.jpg",
                f"{label} 第 {number} 页（扫描页）",
            )
    finally:
        doc.close()


# --------------------------------------------------------------- Excel


def _rels(zf: zipfile.ZipFile, part: str) -> dict[str, str]:
    """Relationship id → target path for an OOXML part."""
    folder, name = posixpath.split(part)
    rels_path = posixpath.join(folder, "_rels", f"{name}.rels")
    if rels_path not in zf.namelist():
        return {}
    root = ElementTree.fromstring(zf.read(rels_path))
    out = {}
    for rel in root.findall("rel:Relationship", _NS):
        target = rel.get("Target", "")
        if rel.get("TargetMode") == "External":
            continue
        out[rel.get("Id")] = posixpath.normpath(
            target.lstrip("/")
            if target.startswith("/")
            else posixpath.join(folder, target)
        )
    return out


def _files(zf: zipfile.ZipFile, prefix: str) -> list[str]:
    """Files (not folder entries) under ``prefix`` in a zip."""
    return [
        n
        for n in zf.namelist()
        if n.startswith(prefix) and not n.endswith("/")
    ]


def _wps_cell_images(zf: zipfile.ZipFile) -> dict[str, str]:
    """WPS ``DISPIMG`` id → media path, from ``xl/cellimages.xml``."""
    part = "xl/cellimages.xml"
    if part not in zf.namelist():
        return {}
    rels = _rels(zf, part)
    root = ElementTree.fromstring(zf.read(part))
    out = {}
    for pic in root.iter(f"{{{_NS['xdr']}}}pic"):
        name = pic.find("xdr:nvPicPr/xdr:cNvPr", _NS)
        blip = pic.find("xdr:blipFill/a:blip", _NS)
        if name is None or blip is None:
            continue
        target = rels.get(blip.get(f"{{{_NS['r']}}}embed"))
        if target:
            out[name.get("name")] = target
    return out


def _ole_package(data: bytes) -> tuple[str, bytes] | None:
    """The file inside an OLE ``Package`` object (Excel's embedded files)."""
    import olefile

    try:
        ole = olefile.OleFileIO(io.BytesIO(data))
    except OSError:
        return None
    try:
        if not ole.exists("\x01Ole10Native"):
            return None
        raw = ole.openstream("\x01Ole10Native").read()
    finally:
        ole.close()
    # size(4) flags(2) label\0 path\0 unknown(4) tmp-path-len(4) tmp-path data-len(4) data
    pos = 6
    end = raw.index(b"\0", pos)
    label = raw[pos:end]
    pos = raw.index(b"\0", end + 1) + 1 + 4
    tmp_len = struct.unpack("<I", raw[pos : pos + 4])[0]
    pos += 4 + tmp_len
    size = struct.unpack("<I", raw[pos : pos + 4])[0]
    payload = raw[pos + 4 : pos + 4 + size]
    for encoding in ("utf-8", "gbk", "latin-1"):
        try:
            return label.decode(encoding), payload
        except UnicodeDecodeError:
            continue
    return "attachment", payload


def _parse_xlsx(data: bytes, label: str, res: _Result, depth: int) -> None:
    import openpyxl
    from openpyxl.utils import get_column_letter

    zf = zipfile.ZipFile(io.BytesIO(data))
    values = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    formulas = openpyxl.load_workbook(io.BytesIO(data), data_only=False)
    cell_images = _wps_cell_images(zf)
    used_media: set[str] = set()
    for ws, wf in zip(values.worksheets, formulas.worksheets):
        res.sheets += 1
        res.lines.append(f"\n### {label} 工作表「{ws.title}」\n")
        rows, picture_cells = [], []
        for row_v, row_f in zip(ws.iter_rows(), wf.iter_rows()):
            cells = []
            for cv, cf in zip(row_v, row_f):
                value = cv.value
                match = _DISPIMG.search(str(cf.value or ""))
                if match:
                    picture_cells.append((cv.coordinate, match.group(1)))
                    value = "[图片]"
                cells.append(value)
            rows.append(cells)
        last_col = max(
            (
                i + 1
                for r in rows
                for i, v in enumerate(r)
                if v not in (None, "")
            ),
            default=0,
        )
        grid = [["行"] + [get_column_letter(i + 1) for i in range(last_col)]]
        grid += [
            [str(n)] + r[:last_col]
            for n, r in enumerate(rows, start=1)
            if any(v not in (None, "") for v in r[:last_col])
        ]
        table = _md_table(grid) if last_col else ""
        if table:
            res.tables += 1
            res.lines.append(
                "位置用「工作表!列行」表示，如 Sheet1!B5。\n\n" + table
            )
        else:
            res.lines.append("（空表）")
        for coordinate, image_id in picture_cells:
            media = cell_images.get(image_id)
            if media and media in zf.namelist():
                used_media.add(media)
                res.add_image(
                    zf.read(media),
                    posixpath.basename(media),
                    f"{label} {ws.title}!{coordinate}",
                )
        for image in getattr(ws, "_images", []):
            try:
                anchor = image.anchor._from
                where = f"{ws.title}!{get_column_letter(anchor.col + 1)}{anchor.row + 1}"
                res.add_image(
                    image._data(),
                    "picture.png",
                    f"{label} {where}（浮动图片）",
                )
            except Exception as exc:  # noqa: BLE001
                res.errors.append(
                    f"{label} {ws.title}：一张浮动图片读不了：{exc}"
                )
    for name in sorted(_files(zf, "xl/embeddings/")):
        if name.endswith(".bin"):
            package = _ole_package(zf.read(name))
            if package:
                _attachment(
                    package[0], package[1], f"{label} 嵌入附件", res, depth
                )
        else:  # embedded .docx/.xlsx/.pdf stored as-is
            _attachment(
                posixpath.basename(name),
                zf.read(name),
                f"{label} 嵌入附件",
                res,
                depth,
            )
    leftovers = [
        n
        for n in _files(zf, "xl/media/")
        if n not in used_media
        and Path(n).suffix.lower() not in {".emf", ".wmf"}
    ]
    for media in leftovers:
        # Floating pictures were added above from openpyxl; repeats are skipped by hash.
        res.add_image(
            zf.read(media),
            posixpath.basename(media),
            f"{label} 未定位图片 {posixpath.basename(media)}",
        )


# ---------------------------------------------------- Word / PowerPoint


def _parse_office(
    data: bytes, suffix: str, label: str, res: _Result, depth: int
) -> None:
    text = _markitdown(data, suffix, label, res)
    if text:
        res.lines.append(f"\n### {label}\n\n{text}")
    media_prefix = {".docx": "word/media/", ".pptx": "ppt/media/"}[suffix]
    embed_prefix = {".docx": "word/embeddings/", ".pptx": "ppt/embeddings/"}[
        suffix
    ]
    zf = zipfile.ZipFile(io.BytesIO(data))
    for name in sorted(_files(zf, "")):
        if name.startswith(media_prefix):
            res.add_image(
                zf.read(name),
                posixpath.basename(name),
                f"{label} 图片 {posixpath.basename(name)}",
            )
        elif name.startswith(embed_prefix):
            if name.endswith(".bin"):
                package = _ole_package(zf.read(name))
                if package:
                    _attachment(
                        package[0], package[1], f"{label} 嵌入附件", res, depth
                    )
            else:
                _attachment(
                    posixpath.basename(name),
                    zf.read(name),
                    f"{label} 嵌入附件",
                    res,
                    depth,
                )


def _markitdown(data: bytes, suffix: str, label: str, res: _Result) -> str:
    try:
        from markitdown import MarkItDown

        return (
            MarkItDown()
            .convert_stream(io.BytesIO(data), file_extension=suffix)
            .text_content.strip()
        )
    except Exception as exc:  # noqa: BLE001
        res.errors.append(f"{label}：转文字失败：{exc}")
        return ""


# ------------------------------------------------------------- dispatch


def _attachment(
    name: str, data: bytes, where: str, res: _Result, depth: int
) -> None:
    if depth >= MAX_ATTACHMENT_DEPTH:
        res.errors.append(f"{where}「{name}」嵌套太深，没有展开")
        return
    target = (
        res.out
        / ATTACHMENT_FOLDER
        / (_safe(Path(name).stem) + Path(name).suffix.lower())
    )
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_bytes(data)
    label = f"{where}「{Path(name).name}」"
    entry = {
        "name": Path(name).name,
        "file": str(target.relative_to(res.out)),
        "bytes": len(data),
    }
    res.attachments.append(entry)
    res.lines.append(f"\n## {label}\n")
    pages_before = res.pages
    _parse_bytes(data, Path(name).suffix.lower(), label, res, depth + 1)
    entry["pages"] = res.pages - pages_before


def _parse_bytes(
    data: bytes, suffix: str, label: str, res: _Result, depth: int
) -> None:
    try:
        if suffix == ".pdf":
            _parse_pdf(data, label, res)
        elif suffix in {".xlsx", ".xlsm"}:
            _parse_xlsx(data, label, res, depth)
        elif suffix in {".docx", ".pptx"}:
            _parse_office(data, suffix, label, res, depth)
        elif suffix in IMAGE_SUFFIXES:
            res.add_image(data, f"{label}{suffix}", f"{label}（图片文件）")
            res.lines.append(f"（{label} 是图片，内容见「需要看图」）")
        elif suffix in TEXT_SUFFIXES:
            res.lines.append(data.decode("utf-8", errors="replace").strip())
        else:
            text = _markitdown(data, suffix, label, res)
            res.lines.append(text or f"（{label}：这种格式读不出文字）")
    except Exception as exc:  # noqa: BLE001 — report, don't fail the upload
        logger.warning("Parsing %s failed", label, exc_info=True)
        res.errors.append(f"{label}：解析出错：{exc}")


def describe(found: dict) -> str:
    """What a file holds, e.g. 「1 个工作表，嵌入附件 7 个（共 8 页）」."""
    parts = []
    if found.get("pages"):
        parts.append(f"{found['pages']} 页")
    if found.get("sheets"):
        parts.append(f"{found['sheets']} 个工作表")
    if found.get("tables"):
        parts.append(f"{found['tables']} 张表格")
    attachments = found.get("attachments") or []
    if attachments:
        pages = sum(a.get("pages", 0) for a in attachments)
        parts.append(
            f"嵌入附件 {len(attachments)} 个"
            + (f"（共 {pages} 页）" if pages else "")
        )
    parts.append(f"需要看图 {found.get('images_to_look', 0)} 张")
    return "，".join(parts)


def parse_folder(raw_folder: Path) -> Path:
    return raw_folder.parent / PARSE_FOLDER


def output_folder(source: Path) -> Path:
    return parse_folder(source.parent) / source.name


def manifest(source: Path) -> dict | None:
    path = output_folder(source) / MANIFEST_FILE
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return None


def is_current(source: Path) -> bool:
    found = manifest(source)
    stat = source.stat()
    return bool(
        found
        and found.get("parser_version") == PARSER_VERSION
        and found.get("source_size") == stat.st_size
        and found.get("source_mtime") == int(stat.st_mtime)
    )


def parse_source(source: Path) -> dict:
    """Parse one file in ``原始资料/`` and return its manifest."""
    out = output_folder(source)
    if out.exists():
        shutil.rmtree(out)
    out.mkdir(parents=True)
    res = _Result(out=out)
    data = source.read_bytes()
    _parse_bytes(data, source.suffix.lower(), source.name, res, depth=0)

    look = [i for i in res.images if i.get("look")]
    own_pages = res.pages - sum(a.get("pages", 0) for a in res.attachments)
    stat = source.stat()
    summary = {
        "source": source.name,
        "source_size": stat.st_size,
        "source_mtime": int(stat.st_mtime),
        "parser_version": PARSER_VERSION,
        "parsed_at": datetime.now().isoformat(timespec="seconds"),
        "pages": own_pages,
        "sheets": res.sheets,
        "tables": res.tables,
        "images_to_look": len(look),
        "images_skipped": len(res.images) - len(look),
        "attachments": res.attachments,
        "images": res.images,
        "errors": res.errors,
        "content_chars": 0,
    }
    head = [
        f"# 解析结果：{source.name}",
        "",
        "程序自动从原件转出的文字，供清洗时阅读；有出入以原件为准。"
        "引用时写「文件名 + 位置」（页码、工作表!单元格、附件名）。",
        "",
        f"- {describe(summary)}"
        f"（另有 {len(res.images) - len(look)} 张图标或装饰已跳过）",
    ]
    if res.errors:
        head.append("- 解析问题：" + "；".join(res.errors))
    tail = []
    if look:
        tail = [
            "",
            "## 需要看图",
            "",
            "这些图片里的文字和信息没有转成文字，要用看图工具读：",
            "",
        ]
        tail += [f"- `{i['file']}` — {i['where']}" for i in look]
    text = _fix_radicals("\n".join(head + res.lines + tail).strip()) + "\n"
    (out / CONTENT_FILE).write_text(text, encoding="utf-8")
    summary["content_chars"] = len(text)
    (out / MANIFEST_FILE).write_text(
        json.dumps(summary, ensure_ascii=False, indent=1), encoding="utf-8"
    )
    return summary


def ensure_parsed(
    raw_folder: Path, names: list[str] | None = None
) -> dict[str, dict]:
    """Parse the files in ``原始资料/`` that have no current result yet.

    Returns name → manifest for the requested files (all files when
    ``names`` is None). Never raises for a single bad file.
    """
    out: dict[str, dict] = {}
    if not raw_folder.is_dir():
        return out
    for source in sorted(raw_folder.iterdir()):
        if not source.is_file() or source.name.startswith("."):
            continue
        if names is not None and source.name not in names:
            continue
        try:
            out[source.name] = (
                manifest(source)
                if is_current(source)
                else parse_source(source)
            )
        except Exception as exc:  # noqa: BLE001
            logger.warning("Parsing %s failed", source, exc_info=True)
            out[source.name] = {
                "source": source.name,
                "errors": [f"解析失败：{exc}"],
            }
    return out
