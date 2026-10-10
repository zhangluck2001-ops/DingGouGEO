# ========= Copyright 2025-2026 @ Eigent.ai All Rights Reserved. =========
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.
# ========= Copyright 2025-2026 @ Eigent.ai All Rights Reserved. =========

import io
import json
import os

import openpyxl
import pytest
from openpyxl.drawing.image import Image as XlImage
from PIL import Image
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from app.geo import source_parse


def _png(size: tuple[int, int]) -> bytes:
    buffer = io.BytesIO()
    # Noise keeps the file above the "too small to matter" byte threshold.
    Image.frombytes("RGB", size, os.urandom(size[0] * size[1] * 3)).save(
        buffer, "PNG"
    )
    return buffer.getvalue()


def _pdf(pages: list[str]) -> bytes:
    buffer = io.BytesIO()
    doc = canvas.Canvas(buffer, pagesize=A4)
    for text in pages:
        if text:
            doc.drawString(72, 760, text)
        doc.showPage()
    doc.save()
    return buffer.getvalue()


@pytest.fixture
def raw(tmp_path):
    folder = tmp_path / "01_品牌知识库" / "原始资料"
    folder.mkdir(parents=True)
    return folder


def test_pdf_text_is_kept_by_page_and_blank_pages_become_images(raw):
    (raw / "label.pdf").write_bytes(
        _pdf(["Lactose, whey protein, sunflower oil and more text here", ""])
    )

    found = source_parse.ensure_parsed(raw)["label.pdf"]

    out = raw.parent / source_parse.PARSE_FOLDER / "label.pdf"
    text = (out / source_parse.CONTENT_FILE).read_text(encoding="utf-8")
    assert found["pages"] == 2
    assert "### label.pdf 第 1 页" in text
    assert "Lactose, whey protein" in text
    assert found["images_to_look"] == 1
    scanned = [i for i in found["images"] if i.get("look")][0]
    assert scanned["where"] == "label.pdf 第 2 页（扫描页）"
    assert (out / scanned["file"]).is_file()
    assert f"`{scanned['file']}` — label.pdf 第 2 页（扫描页）" in text


def test_excel_cells_keep_their_address_and_pictures_their_cell(raw):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "产品"
    ws["A1"], ws["B1"] = "字段", "内容"
    ws["A2"], ws["B2"] = "品牌发源地", "英国"
    picture = XlImage(io.BytesIO(_png((400, 300))))
    ws.add_image(picture, "C2")
    ws.add_image(XlImage(io.BytesIO(_png((40, 40)))), "D2")  # an icon
    wb.save(raw / "brand.xlsx")

    found = source_parse.ensure_parsed(raw)["brand.xlsx"]

    text = (
        source_parse.output_folder(raw / "brand.xlsx")
        / source_parse.CONTENT_FILE
    ).read_text(encoding="utf-8")
    assert "| 2 | 品牌发源地 | 英国 |" in text
    assert found["sheets"] == 1
    assert found["images_to_look"] == 1
    assert found["images_skipped"] == 1
    look = [i for i in found["images"] if i.get("look")][0]
    assert look["where"] == "brand.xlsx 产品!C2（浮动图片）"


def test_cjk_radicals_from_pdfs_are_read_as_normal_characters():
    assert source_parse._fix_radicals("部分⽔解乳清蛋⽩（⽜乳）") == (
        "部分水解乳清蛋白（牛乳）"
    )


def test_a_file_is_parsed_once_until_it_changes(raw):
    source = raw / "notes.md"
    source.write_text("SMA 1915 年", encoding="utf-8")
    first = source_parse.ensure_parsed(raw)["notes.md"]

    again = source_parse.ensure_parsed(raw)["notes.md"]
    assert again["parsed_at"] == first["parsed_at"]

    source.write_text("SMA 1915 年，英国", encoding="utf-8")
    os.utime(source, (source.stat().st_atime, source.stat().st_mtime + 5))
    changed = source_parse.ensure_parsed(raw)["notes.md"]
    assert changed["source_size"] != first["source_size"]
    assert "英国" in (
        source_parse.output_folder(source) / source_parse.CONTENT_FILE
    ).read_text(encoding="utf-8")


def test_a_broken_file_is_reported_not_raised(raw):
    (raw / "broken.xlsx").write_bytes(b"not a workbook")

    found = source_parse.ensure_parsed(raw)["broken.xlsx"]

    assert found["errors"]
    manifest = json.loads(
        (
            source_parse.output_folder(raw / "broken.xlsx")
            / source_parse.MANIFEST_FILE
        ).read_text(encoding="utf-8")
    )
    assert manifest["errors"] == found["errors"]


def test_attachment_pages_are_not_counted_as_the_files_own():
    found = {
        "pages": 0,
        "sheets": 1,
        "tables": 1,
        "attachments": [{"pages": 2}, {"pages": 6}],
        "images_to_look": 3,
    }
    assert source_parse.describe(found) == (
        "1 个工作表，1 张表格，嵌入附件 2 个（共 8 页），需要看图 3 张"
    )
